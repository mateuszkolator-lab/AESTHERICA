from fastapi import APIRouter, Depends, HTTPException, Header
from fastapi.responses import StreamingResponse
from typing import Optional, List, Any
from pydantic import BaseModel
import uuid
import re
import ast
import io
import operator
from datetime import datetime, timezone, timedelta

import openpyxl
from openpyxl.utils import get_column_letter

from models.database import get_db
from routers.auth import verify_token, get_patient_location_filter
from routers.audit import log_patient_change
from routers.survey_presets import PRESET_TEMPLATES

router = APIRouter(prefix="/surveys", tags=["Surveys / PROMs"])

def get_auth(authorization: str = Header(None)):
    return verify_token(authorization)

def now_iso():
    return datetime.now(timezone.utc).isoformat()

# ---------- Schemas ----------
class TemplateIn(BaseModel):
    name: str
    code: Optional[str] = None
    description: Optional[str] = None
    instructions: Optional[str] = None
    active: bool = True
    timepoints: List[dict] = []      # [{key, label}]
    sections: List[dict] = []        # [{key, label}]
    questions: List[dict] = []       # see compute_scores
    scoring: dict = {}               # {mode, multiplier, scale_to_100, formula, subscales:[...]}

class SyncAssignments(BaseModel):
    template_ids: List[str]

class PublicSubmit(BaseModel):
    answers: dict
    filled_via: Optional[str] = "link"

# ---------- Scoring engine ----------
SAFE_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
    ast.Div: operator.truediv, ast.Pow: operator.pow, ast.USub: operator.neg, ast.UAdd: operator.pos,
}
SAFE_FUNCS = {"sum": sum, "min": min, "max": max, "abs": abs, "round": round, "len": len,
              "mean": lambda xs: (sum(xs) / len(xs)) if xs else 0}

def safe_eval(expr: str, variables: dict):
    def _eval(node):
        if isinstance(node, ast.Expression):
            return _eval(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name):
            if node.id in variables:
                return variables[node.id]
            raise ValueError(f"Nieznana zmienna: {node.id}")
        if isinstance(node, ast.BinOp) and type(node.op) in SAFE_OPS:
            return SAFE_OPS[type(node.op)](_eval(node.left), _eval(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in SAFE_OPS:
            return SAFE_OPS[type(node.op)](_eval(node.operand))
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in SAFE_FUNCS:
            return SAFE_FUNCS[node.func.id](*[_eval(a) for a in node.args])
        if isinstance(node, (ast.List, ast.Tuple)):
            return [_eval(e) for e in node.elts]
        raise ValueError("Niedozwolone wyrażenie")
    return _eval(ast.parse(expr, mode="eval"))

def question_score(q: dict, value: Any):
    """Returns (score, max_score) for a single question or (None, None) if not scorable."""
    qtype = q.get("type")
    if value is None or value == "" or qtype in ("text",):
        return None, None
    if qtype in ("scale", "single", "yesno"):
        options = q.get("options") or []
        if options:
            scores = [float(o.get("score", o.get("value", 0)) or 0) for o in options]
            chosen = next((o for o in options if str(o.get("value")) == str(value)), None)
            if chosen is None:
                return None, None
            score = float(chosen.get("score", chosen.get("value", 0)) or 0)
            lo, hi = min(scores), max(scores)
        else:
            lo, hi = float(q.get("min", 0)), float(q.get("max", 10))
            try:
                score = float(value)
            except (TypeError, ValueError):
                return None, None
        if q.get("reverse"):
            score = lo + hi - score
        return score, hi
    if qtype == "multi":
        options = q.get("options") or []
        vals = value if isinstance(value, list) else [value]
        score = sum(float(o.get("score", 0) or 0) for o in options if str(o.get("value")) in [str(v) for v in vals])
        return score, sum(max(0.0, float(o.get("score", 0) or 0)) for o in options)
    if qtype == "number":
        try:
            return float(value), float(q.get("max", value))
        except (TypeError, ValueError):
            return None, None
    return None, None

def aggregate(scores: List[float], maxes: List[float], cfg: dict):
    if not scores:
        return None
    mode = cfg.get("mode", "sum")
    total = sum(scores) / len(scores) if mode == "mean" else sum(scores)
    max_total = (sum(maxes) / len(maxes)) if mode == "mean" else sum(maxes)
    mult = cfg.get("multiplier")
    if mult not in (None, "", 0):
        total *= float(mult)
        max_total *= float(mult)
    if cfg.get("scale_to_100") and max_total:
        total = total / max_total * 100
        max_total = 100
    return {"value": round(total, 2), "max": round(max_total, 2)}

def compute_scores(template: dict, answers: dict) -> dict:
    per_q = {}
    for q in template.get("questions", []):
        s, m = question_score(q, answers.get(q["id"]))
        if s is not None:
            per_q[q["id"]] = (s, m)
    scoring = template.get("scoring") or {}
    result = {"per_question": {k: v[0] for k, v in per_q.items()}}

    all_scores = [v[0] for v in per_q.values()]
    all_max = [v[1] for v in per_q.values()]
    result["total"] = aggregate(all_scores, all_max, scoring)

    subs = {}
    for sub in scoring.get("subscales") or []:
        ids = sub.get("question_ids") or [q["id"] for q in template.get("questions", []) if q.get("section") == sub.get("section")]
        sc = [per_q[i][0] for i in ids if i in per_q]
        mx = [per_q[i][1] for i in ids if i in per_q]
        subs[sub.get("key") or sub.get("label")] = {"label": sub.get("label"), **(aggregate(sc, mx, sub) or {"value": None, "max": None})}
    result["subscales"] = subs

    formula = (scoring.get("formula") or "").strip()
    if formula:
        variables = {f"q{i+1}": per_q.get(q["id"], (0, 0))[0] for i, q in enumerate(template.get("questions", []))}
        variables["total"] = result["total"]["value"] if result["total"] else 0
        variables["items"] = all_scores
        for k, v in subs.items():
            variables[k] = v.get("value") or 0
        try:
            result["total"] = {"value": round(float(safe_eval(formula, variables)), 2), "max": result["total"]["max"] if result["total"] else None, "formula": formula}
        except Exception as e:
            result["formula_error"] = str(e)
    return result

# ---------- Seed ----------
KEY_OFFSETS = {"pre": 0, "before": 0, "1w": 7, "2w": 14, "1m": 30, "3m": 90, "6m": 180, "12m": 365, "1y": 365, "24m": 730, "2y": 730}

def infer_offset(key: str):
    if key in KEY_OFFSETS:
        return KEY_OFFSETS[key]
    m = re.match(r"^(\d+)(d|w|m|y)$", key or "")
    if m:
        n, unit = int(m.group(1)), m.group(2)
        return n * {"d": 1, "w": 7, "m": 30, "y": 365}[unit]
    return None

async def seed_survey_templates():
    db = get_db()
    for preset in PRESET_TEMPLATES:
        if not await db.survey_templates.find_one({"code": preset["code"]}):
            await db.survey_templates.insert_one({**preset, "id": str(uuid.uuid4()), "preset": True, "created_at": now_iso(), "updated_at": now_iso()})
    # Migration: infer offset_days for timepoints missing it
    async for t in db.survey_templates.find({"timepoints.offset_days": {"$exists": False}}):
        tps = t.get("timepoints", [])
        changed = False
        for tp in tps:
            if "offset_days" not in tp:
                tp["offset_days"] = infer_offset(tp.get("key"))
                changed = True
        if changed:
            await db.survey_templates.update_one({"_id": t["_id"]}, {"$set": {"timepoints": tps}})

# ---------- Templates ----------
def _normalize_template(data: dict) -> dict:
    for q in data.get("questions", []):
        q.setdefault("id", str(uuid.uuid4())[:8])
    if not data.get("timepoints"):
        data["timepoints"] = [{"key": "pre", "label": "Przed zabiegiem"}]
    return data

@router.get("/templates")
async def list_templates(include_inactive: bool = False, user: dict = Depends(get_auth)):
    db = get_db()
    q = {} if include_inactive else {"active": {"$ne": False}}
    templates = await db.survey_templates.find(q, {"_id": 0}).sort("name", 1).to_list(200)
    return templates

@router.post("/templates")
async def create_template(data: TemplateIn, user: dict = Depends(get_auth)):
    db = get_db()
    doc = _normalize_template(data.model_dump())
    doc.update({"id": str(uuid.uuid4()), "created_at": now_iso(), "updated_at": now_iso()})
    await db.survey_templates.insert_one(doc)
    return {"id": doc["id"]}

@router.get("/templates/{template_id}")
async def get_template(template_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    t = await db.survey_templates.find_one({"id": template_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")
    return t

@router.put("/templates/{template_id}")
async def update_template(template_id: str, data: TemplateIn, user: dict = Depends(get_auth)):
    db = get_db()
    doc = _normalize_template(data.model_dump())
    doc["updated_at"] = now_iso()
    res = await db.survey_templates.update_one({"id": template_id}, {"$set": doc})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Template not found")
    # Refresh pending assignments' timepoints (add new ones for assigned patients)
    patient_ids = await db.patients.distinct("id", {"survey_template_ids": template_id})
    for pid in patient_ids:
        await _sync_patient_assignments(db, pid, None, only_template=template_id)
    return {"message": "updated"}

@router.delete("/templates/{template_id}")
async def delete_template(template_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    completed = await db.survey_assignments.count_documents({"template_id": template_id, "status": "completed"})
    if completed:
        await db.survey_templates.update_one({"id": template_id}, {"$set": {"active": False, "updated_at": now_iso()}})
        return {"message": "deactivated", "completed_responses": completed}
    await db.survey_templates.delete_one({"id": template_id})
    await db.survey_assignments.delete_many({"template_id": template_id})
    return {"message": "deleted"}

@router.post("/templates/{template_id}/duplicate")
async def duplicate_template(template_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    t = await db.survey_templates.find_one({"id": template_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")
    t.update({"id": str(uuid.uuid4()), "name": f"{t['name']} (kopia)", "code": None, "preset": False,
              "created_at": now_iso(), "updated_at": now_iso()})
    await db.survey_templates.insert_one(t)
    return {"id": t["id"]}

@router.post("/templates/preview-score")
async def preview_score(payload: dict, user: dict = Depends(get_auth)):
    """Compute score for given template definition + answers (builder test)."""
    return compute_scores(payload.get("template", {}), payload.get("answers", {}))

# ---------- Assignments ----------
async def _sync_patient_assignments(db, patient_id: str, template_ids: Optional[List[str]], only_template: str = None):
    patient = await db.patients.find_one({"id": patient_id}, {"_id": 0})
    if not patient:
        raise HTTPException(status_code=404, detail="Patient not found")
    if template_ids is None:
        template_ids = patient.get("survey_template_ids") or []
    else:
        await db.patients.update_one({"id": patient_id}, {"$set": {"survey_template_ids": template_ids}})

    existing = await db.survey_assignments.find({"patient_id": patient_id}, {"_id": 0}).to_list(500)
    existing_keys = {(a["template_id"], a["timepoint_key"]) for a in existing}

    templates = await db.survey_templates.find({"id": {"$in": template_ids}}, {"_id": 0}).to_list(100)
    for t in templates:
        if only_template and t["id"] != only_template:
            continue
        for tp in t.get("timepoints", []):
            if (t["id"], tp["key"]) not in existing_keys:
                await db.survey_assignments.insert_one({
                    "id": str(uuid.uuid4()), "token": uuid.uuid4().hex,
                    "patient_id": patient_id, "template_id": t["id"],
                    "timepoint_key": tp["key"], "timepoint_label": tp.get("label"),
                    "status": "pending", "created_at": now_iso()
                })
    # Remove pending assignments for unselected templates
    if not only_template:
        await db.survey_assignments.delete_many({"patient_id": patient_id, "status": "pending", "template_id": {"$nin": template_ids}})

@router.get("/patients/{patient_id}")
async def get_patient_assignments(patient_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    assignments = await db.survey_assignments.find({"patient_id": patient_id}, {"_id": 0}).to_list(500)
    templates = {t["id"]: t for t in await db.survey_templates.find({}, {"_id": 0}).to_list(200)}
    for a in assignments:
        t = templates.get(a["template_id"], {})
        a["template_name"] = t.get("name", "?")
        a["template_code"] = t.get("code")
        order = [tp["key"] for tp in t.get("timepoints", [])]
        a["_order"] = order.index(a["timepoint_key"]) if a["timepoint_key"] in order else 99
    assignments.sort(key=lambda a: (a["template_name"], a["_order"]))
    for a in assignments:
        a.pop("_order", None)
    return assignments

@router.post("/patients/{patient_id}/sync")
async def sync_patient_assignments(patient_id: str, data: SyncAssignments, user: dict = Depends(get_auth)):
    db = get_db()
    await _sync_patient_assignments(db, patient_id, data.template_ids)
    return {"message": "synced"}

@router.post("/assignments/{assignment_id}/reset")
async def reset_assignment(assignment_id: str, user: dict = Depends(get_auth)):
    """Allow re-filling: archive previous answers into history."""
    db = get_db()
    a = await db.survey_assignments.find_one({"id": assignment_id}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Assignment not found")
    update = {"$set": {"status": "pending", "token": uuid.uuid4().hex}, "$unset": {"answers": "", "score": "", "completed_at": "", "filled_via": ""}}
    if a.get("status") == "completed":
        update["$push"] = {"history": {"answers": a.get("answers"), "score": a.get("score"), "completed_at": a.get("completed_at")}}
    await db.survey_assignments.update_one({"id": assignment_id}, update)
    return {"message": "reset"}

@router.delete("/assignments/{assignment_id}")
async def delete_assignment(assignment_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    await db.survey_assignments.delete_one({"id": assignment_id})
    return {"message": "deleted"}

@router.get("/assignments/{assignment_id}")
async def get_assignment(assignment_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    a = await db.survey_assignments.find_one({"id": assignment_id}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Assignment not found")
    a["template"] = await db.survey_templates.find_one({"id": a["template_id"]}, {"_id": 0})
    a["patient"] = await db.patients.find_one({"id": a["patient_id"]}, {"_id": 0, "first_name": 1, "last_name": 1, "id": 1})
    return a

# ---------- Public (patient) ----------
@router.get("/public/{token}")
async def public_get(token: str):
    db = get_db()
    a = await db.survey_assignments.find_one({"token": token}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Ankieta nie istnieje lub link wygasł")
    t = await db.survey_templates.find_one({"id": a["template_id"]}, {"_id": 0, "scoring": 0})
    p = await db.patients.find_one({"id": a["patient_id"]}, {"_id": 0, "first_name": 1})
    return {
        "status": a["status"], "timepoint_label": a.get("timepoint_label"),
        "completed_at": a.get("completed_at"),
        "patient_first_name": (p or {}).get("first_name"),
        "template": t
    }

@router.post("/public/{token}")
async def public_submit(token: str, data: PublicSubmit):
    db = get_db()
    a = await db.survey_assignments.find_one({"token": token}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Ankieta nie istnieje")
    if a["status"] == "completed":
        raise HTTPException(status_code=409, detail="Ta ankieta została już wypełniona")
    t = await db.survey_templates.find_one({"id": a["template_id"]}, {"_id": 0})
    missing = [q["id"] for q in t.get("questions", []) if q.get("required", True) and data.answers.get(q["id"]) in (None, "", [])]
    if missing:
        raise HTTPException(status_code=422, detail={"missing": missing})
    score = compute_scores(t, data.answers)
    await db.survey_assignments.update_one({"token": token}, {"$set": {
        "status": "completed", "answers": data.answers, "score": score,
        "completed_at": now_iso(), "filled_via": data.filled_via or "link"
    }})
    await log_patient_change(patient_id=a["patient_id"], action="survey_completed", user={"user_id": None, "user_name": "Pacjent"},
                             details=f"Pacjent wypełnił ankietę {t.get('name')} ({a.get('timepoint_label')})")
    return {"message": "saved", "score": score}

# ---------- Reminders ----------
PRE_WINDOW_DAYS = 14   # show "before surgery" surveys from 14 days ahead
POST_WINDOW_DAYS = 7   # show post-op surveys from 7 days before due date

@router.get("/due")
async def due_surveys(user: dict = Depends(get_auth)):
    """Pending assignments whose timepoint is due (or overdue) based on patient's surgery date."""
    db = get_db()
    today = datetime.now(timezone.utc).date()
    pending = await db.survey_assignments.find({"status": "pending"}, {"_id": 0}).to_list(20000)
    if not pending:
        return []
    pq = {"id": {"$in": list({a["patient_id"] for a in pending})}, "surgery_date": {"$nin": [None, ""]},
          "status": {"$nin": ["cancelled"]}}
    loc_filter = get_patient_location_filter(user)
    if loc_filter:
        pq = {"$and": [pq, loc_filter]}
    patients = {p["id"]: p for p in await db.patients.find(pq, {"_id": 0, "visits": 0, "photos": 0}).to_list(5000)}
    templates = {t["id"]: t for t in await db.survey_templates.find({}, {"_id": 0}).to_list(200)}
    out = []
    for a in pending:
        p = patients.get(a["patient_id"])
        t = templates.get(a["template_id"])
        if not p or not t:
            continue
        tp = next((x for x in t.get("timepoints", []) if x["key"] == a["timepoint_key"]), None)
        if not tp or tp.get("offset_days") is None:
            continue
        try:
            sd = datetime.strptime(p["surgery_date"][:10], "%Y-%m-%d").date()
        except ValueError:
            continue
        offset = int(tp["offset_days"])
        due = sd + timedelta(days=offset)
        window = PRE_WINDOW_DAYS if offset <= 0 else POST_WINDOW_DAYS
        if (due - today).days > window:
            continue
        if offset <= 0 and today > sd:
            continue  # pre-op survey no longer relevant after surgery
        out.append({
            "assignment_id": a["id"], "token": a["token"],
            "patient_id": p["id"], "first_name": p.get("first_name"), "last_name": p.get("last_name"),
            "phone": p.get("phone"), "email": p.get("email"), "location_id": p.get("location_id"),
            "surgery_date": p["surgery_date"], "template_name": t.get("name"), "template_code": t.get("code"),
            "timepoint_label": a.get("timepoint_label") or tp.get("label"),
            "due_date": due.isoformat(), "days_overdue": (today - due).days
        })
    out.sort(key=lambda x: -x["days_overdue"])
    return out

# ---------- Results ----------
async def _results_rows(db, user, template_id: str, timepoint: Optional[str] = None):
    q = {"template_id": template_id, "status": "completed"}
    if timepoint:
        q["timepoint_key"] = timepoint
    assignments = await db.survey_assignments.find(q, {"_id": 0}).sort("completed_at", -1).to_list(5000)
    pq = {"id": {"$in": list({a["patient_id"] for a in assignments})}}
    loc_filter = get_patient_location_filter(user)
    if loc_filter:
        pq = {"$and": [pq, loc_filter]}
    patients = {p["id"]: p for p in await db.patients.find(pq, {"_id": 0, "visits": 0}).to_list(5000)}
    rows = []
    for a in assignments:
        p = patients.get(a["patient_id"])
        if not p:
            continue
        rows.append({**a, "patient": {"id": p["id"], "first_name": p.get("first_name"), "last_name": p.get("last_name"),
                                     "procedure_type": p.get("procedure_type"), "surgery_date": p.get("surgery_date"),
                                     "location_id": p.get("location_id")}})
    return rows

@router.get("/results")
async def get_results(template_id: str, timepoint: Optional[str] = None, user: dict = Depends(get_auth)):
    db = get_db()
    return await _results_rows(db, user, template_id, timepoint)

@router.get("/results/summary")
async def results_summary(user: dict = Depends(get_auth)):
    db = get_db()
    pipeline = [{"$group": {"_id": {"t": "$template_id", "s": "$status"}, "n": {"$sum": 1}}}]
    agg = await db.survey_assignments.aggregate(pipeline).to_list(500)
    out = {}
    for r in agg:
        out.setdefault(r["_id"]["t"], {"pending": 0, "completed": 0})[r["_id"]["s"]] = r["n"]
    return out

class BulkExport(BaseModel):
    template_ids: List[str] = []
    timepoints: List[str] = []
    location_id: Optional[str] = None
    procedure_type: Optional[str] = None
    surgery_from: Optional[str] = None
    surgery_to: Optional[str] = None
    statuses: List[str] = []
    patient_ids: List[str] = []

def _answer_label(q: dict, v):
    if q.get("type") in ("scale", "single", "yesno") and q.get("options"):
        opt = next((o for o in q["options"] if str(o.get("value")) == str(v)), None)
        return f"{v} – {opt.get('label')}" if opt and opt.get("label") else v
    if isinstance(v, list):
        return ", ".join(map(str, v))
    return v

def _bold(ws, row, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).font = openpyxl.styles.Font(bold=True)

@router.post("/export-bulk")
async def export_bulk(f: BulkExport, user: dict = Depends(get_auth)):
    """Export several surveys for a filtered group of patients: summary sheet + one sheet per survey."""
    db = get_db()
    # Patient group
    pq = {}
    loc_filter = get_patient_location_filter(user)
    if loc_filter:
        pq = {"$and": [pq, loc_filter]} if pq else dict(loc_filter)
    if f.location_id:
        pq["location_id"] = f.location_id
    if f.procedure_type:
        pq["procedure_type"] = f.procedure_type
    if f.statuses:
        pq["status"] = {"$in": f.statuses}
    if f.surgery_from or f.surgery_to:
        rng = {}
        if f.surgery_from:
            rng["$gte"] = f.surgery_from
        if f.surgery_to:
            rng["$lte"] = f.surgery_to
        pq["surgery_date"] = rng
    if f.patient_ids:
        pq["id"] = {"$in": f.patient_ids}
    patients = await db.patients.find(pq, {"_id": 0, "visits": 0}).sort("last_name", 1).to_list(5000)
    pids = [p["id"] for p in patients]
    pmap = {p["id"]: p for p in patients}

    tq = {"active": {"$ne": False}} if not f.template_ids else {"id": {"$in": f.template_ids}}
    templates = await db.survey_templates.find(tq, {"_id": 0}).sort("name", 1).to_list(100)
    locations = {l["id"]: l["name"] for l in await db.locations.find({}, {"_id": 0}).to_list(100)}

    aq = {"patient_id": {"$in": pids}, "template_id": {"$in": [t["id"] for t in templates]}}
    if f.timepoints:
        aq["timepoint_key"] = {"$in": f.timepoints}
    assignments = await db.survey_assignments.find(aq, {"_id": 0}).to_list(20000)
    by_key = {(a["patient_id"], a["template_id"], a["timepoint_key"]): a for a in assignments}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    head = ["Nazwisko", "Imię", "Status", "Placówka", "Zabieg", "Data zabiegu"]
    cols = []
    for t in templates:
        for tp in t.get("timepoints", []):
            if f.timepoints and tp["key"] not in f.timepoints:
                continue
            cols.append((t, tp))
            head.append(f"{t.get('code') or t['name']} – {tp['label']}")
            for s in (t.get("scoring") or {}).get("subscales") or []:
                head.append(f"{t.get('code') or t['name']} – {tp['label']} – {s.get('label') or s.get('key')}")
    for c, h in enumerate(head, 1):
        ws.cell(row=1, column=c, value=h)
    _bold(ws, 1, len(head))
    status_pl = {"consultation": "Konsultacja", "planned": "Zaplanowany", "awaiting": "Oczekujący", "thinking": "Zastanawia się", "operated": "Zoperowany", "cancelled": "Rezygnacja"}
    r = 2
    for p in patients:
        row = [p.get("last_name"), p.get("first_name"), status_pl.get(p.get("status"), p.get("status")),
               locations.get(p.get("location_id"), ""), p.get("procedure_type"), p.get("surgery_date")]
        for t, tp in cols:
            a = by_key.get((p["id"], t["id"], tp["key"]))
            sc = (a or {}).get("score") or {}
            row.append((sc.get("total") or {}).get("value") if a and a.get("status") == "completed" else ("oczekuje" if a else None))
            for s in (t.get("scoring") or {}).get("subscales") or []:
                k = s.get("key") or s.get("label")
                row.append(((sc.get("subscales") or {}).get(k) or {}).get("value") if a and a.get("status") == "completed" else None)
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        r += 1
    for c in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20

    # Detail sheet per template
    for t in templates:
        wsd = wb.create_sheet(title=(t.get("code") or t["name"])[:30])
        qs = t.get("questions", [])
        sub_keys = [s.get("key") or s.get("label") for s in (t.get("scoring") or {}).get("subscales") or []]
        h = ["Nazwisko", "Imię", "Placówka", "Zabieg", "Data zabiegu", "Moment", "Data wypełnienia", "Wynik", "Max"] + sub_keys + \
            [f"P{i+1}: {q.get('text','')[:60]}" for i, q in enumerate(qs)]
        for c, v in enumerate(h, 1):
            wsd.cell(row=1, column=c, value=v)
        _bold(wsd, 1, len(h))
        rr = 2
        for a in sorted([a for a in assignments if a["template_id"] == t["id"] and a.get("status") == "completed"],
                        key=lambda a: (pmap[a["patient_id"]].get("last_name", ""), a["timepoint_key"])):
            p = pmap[a["patient_id"]]
            sc = a.get("score") or {}
            vals = [p.get("last_name"), p.get("first_name"), locations.get(p.get("location_id"), ""), p.get("procedure_type"), p.get("surgery_date"),
                    a.get("timepoint_label"), (a.get("completed_at") or "")[:16].replace("T", " "),
                    (sc.get("total") or {}).get("value"), (sc.get("total") or {}).get("max")]
            vals += [((sc.get("subscales") or {}).get(k) or {}).get("value") for k in sub_keys]
            vals += [_answer_label(q, (a.get("answers") or {}).get(q["id"])) for q in qs]
            for c, v in enumerate(vals, 1):
                wsd.cell(row=rr, column=c, value=v)
            rr += 1
        for c in range(1, len(h) + 1):
            wsd.column_dimensions[get_column_letter(c)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"ankiety_zbiorczo_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@router.get("/export")
async def export_results(template_id: str, user: dict = Depends(get_auth)):
    db = get_db()
    t = await db.survey_templates.find_one({"id": template_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")
    rows = await _results_rows(db, user, template_id)
    questions = t.get("questions", [])
    sub_keys = [s.get("key") or s.get("label") for s in (t.get("scoring") or {}).get("subscales") or []]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (t.get("code") or t.get("name") or "Ankieta")[:30]
    headers = ["Nazwisko", "Imię", "Zabieg", "Data zabiegu", "Moment", "Data wypełnienia", "Wynik", "Max"] + \
              [f"{k}" for k in sub_keys] + [f"P{i+1}: {q.get('text','')[:60]}" for i, q in enumerate(questions)]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = openpyxl.styles.Font(bold=True)
    for r, row in enumerate(rows, 2):
        p, sc = row["patient"], row.get("score") or {}
        vals = [p["last_name"], p["first_name"], p.get("procedure_type"), p.get("surgery_date"), row.get("timepoint_label"),
                (row.get("completed_at") or "")[:16].replace("T", " "),
                (sc.get("total") or {}).get("value"), (sc.get("total") or {}).get("max")]
        vals += [((sc.get("subscales") or {}).get(k) or {}).get("value") for k in sub_keys]
        answers = row.get("answers") or {}
        for q in questions:
            v = answers.get(q["id"])
            if q.get("type") in ("scale", "single", "yesno") and q.get("options"):
                opt = next((o for o in q["options"] if str(o.get("value")) == str(v)), None)
                v = f"{v} – {opt.get('label')}" if opt and opt.get("label") else v
            elif isinstance(v, list):
                v = ", ".join(map(str, v))
            vals.append(v)
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"ankieta_{(t.get('code') or t.get('name') or 'wyniki').replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})
